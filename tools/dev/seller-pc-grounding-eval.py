#!/usr/bin/env python3
"""How well a review finds the guidance written for its problem — measured on the server, without a model.

Runs against the judging copy only (`SELLER_PC_DEMO=judge tools/dev/seller-pc-demo.sh up`). For the length of the run
it registers every open 안내 제안 (ISSUE_GUIDANCE) as knowledge — a 【확인: …】 blank becomes 「해당 내용」, since a blank
is never retrieved — asks `/reply/draft/grounds` for every review, and removes what it registered on the way out, so
the copy is as it was.

    python3 -I tools/dev/seller-pc-grounding-eval.py [--cutoff 2026-04-22]

Two readings of the same reviews:

  full   — the guidance with the customer phrases it has now. Those phrases were picked by reading these very reviews,
           so this is the ceiling, not the expectation.
  past   — reviews after --cutoff, with each guidance keeping only the phrases already said in a review before it:
           what the phrase list would have found had it been written on the cutoff day. The honest estimate for
           reviews that have not arrived yet.

The reference for 「this review reports this problem」 is each guidance's own `match`/`complaint` pattern in
tools/dev/.run/issue-guidance.json — written by the same hand as the phrases, so it is a check against the lane, not
ground truth. Every attached pair and every miss is written to tools/dev/.run/seller-pc-judge/근거평가.xlsx (gitignored)
for a person to mark 맞다 / 일부 / 아니다. A later run keeps every mark already in the sheet (by guidance and review
text), and once marks exist it also counts by them — the marks, not the patterns, are the reference worth reporting.
Prints counts only: review text stays in the sheet.
"""
import argparse
import importlib.util
import json
import re
import subprocess
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location("guidance", ROOT / "tools/dev/seller-pc-issue-guidance.py")
g = importlib.util.module_from_spec(spec)
spec.loader.exec_module(g)

PORT, DB = 18690, "sellerpc_judge"
OUT = ROOT / "tools/dev/.run/seller-pc-judge/근거평가.xlsx"
FILLER = "해당 내용"


def squash(text):
    return re.sub(r"\s+", "", text or "").lower()


def psql(sql):
    subprocess.run([g.PSQL, "-h", "127.0.0.1", "-p", "5432", "-U", "sellerops", "-d", DB, "-qc", sql], check=True)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cutoff", default="2026-04-22", help="reviews on or after this date are the unseen half")
    args = ap.parse_args()

    guidance = {e["subject"]: e for e in json.loads((ROOT / "tools/dev/.run/issue-guidance.json").read_text())["guidance"]}
    token = g.call(PORT, "POST", "/api/auth/login", {"email": "demo@sellerops.ai", "password": g.seed_password()})["token"]

    open_ = [c for c in g.call(PORT, "GET", "/api/knowledge/candidates", token=token)
             if c["origin"] == "ISSUE_GUIDANCE" and c["subject"] in guidance]
    accepted = {}  # sourceId -> (subject, candidateId, org)
    try:
        for c in open_:
            e = guidance[c["subject"]]
            org = c["productId"] is None
            content = re.sub(r"【확인:[^】]*】", FILLER, c["content"])
            body = {"title": c["subject"], "content": content, "customerPhrases": c.get("customerPhrases") or []}
            body.update({"orgType": "SHIPPING_POLICY"} if org else {"sourceType": "USAGE"})
            acc = g.call(PORT, "POST", f"/api/knowledge/candidates/{c['id']}/accept", body, token=token)
            accepted[acc["sourceId"]] = (c["subject"], c["id"], org)
        print(f"registered {len(accepted)} guidance for the run")

        reviews = g.sql_rows(DB, "select r.id, r.rating, r.received_at::date, coalesce(r.source_product_ref,''), r.body, "
                                 "(select a.id from seller_accounts a where a.org_id=r.org_id and a.channel_id=r.channel_id "
                                 "limit 1) from reviews r where length(trim(r.body)) > 0 order by r.received_at")

        def grounds(rid, account):
            res = g.call(PORT, "GET", f"/api/seller-accounts/{account}/attention/items/review:{rid}/reply/draft/grounds",
                         token=token)
            return {x["sourceId"]: x.get("chunkId") is None for x in res if x.get("sourceId") in accepted}

        def reference(subject, ref, body):
            e = guidance[subject]
            return (not e["refs"] or ref in e["refs"]) and g.reports(e, body)

        full = {rid: grounds(rid, a) for rid, _, _, _, _, a in reviews}
        print(f"read {len(full)} reviews (full phrases)")

        # The phrases each guidance would have had on the cutoff day: those already said in an earlier review.
        before = [squash(b) for _, _, d, _, b, _ in reviews if str(d) < args.cutoff]
        kept = {}
        for sid, (subject, _, org) in accepted.items():
            phrases = guidance[subject]["customer_phrases"]
            seen = [p for p in phrases if any(squash(p) in b for b in before)]
            kept[subject] = (len(seen), len(phrases))
            stored = "\n".join(seen).replace("'", "''")
            table = "org_knowledge_sources" if org else "product_knowledge_sources"
            psql(f"update {table} set customer_phrases = " + ("null" if not seen else f"'{stored}'") + f" where id = '{sid}'")
        after = [r for r in reviews if str(r[2]) >= args.cutoff]
        past = {rid: grounds(rid, a) for rid, _, _, _, _, a in after}
        print(f"read {len(past)} reviews on or after {args.cutoff} (past phrases only)")

        rows = []
        summary = {}
        for subject in sorted({s for s, _, _ in accepted.values()}):
            # One subject can be one guidance per product (the dispenser ones are on two): any of them counts.
            sids = {k for k, v in accepted.items() if v[0] == subject}
            s = summary.setdefault(subject, {"labeled": 0, "found": 0, "wrong": 0, "labeled_after": 0,
                                             "found_after_full": 0, "found_after_past": 0, "wrong_after_past": 0,
                                             "kept": kept.get(subject)})
            for rid, rating, day, ref, body, _ in reviews:
                labeled = reference(subject, ref, body)
                hit = bool(sids & full[rid].keys())
                late = rid in past
                if labeled:
                    s["labeled"] += 1
                    s["found"] += hit
                if hit and not labeled:
                    s["wrong"] += 1
                if late:
                    hit_past = bool(sids & past[rid].keys())
                    if labeled:
                        s["labeled_after"] += 1
                        s["found_after_full"] += hit
                        s["found_after_past"] += hit_past
                    if hit_past and not labeled:
                        s["wrong_after_past"] += 1
                if hit or labeled:
                    rows.append((subject, rid, rating, str(day), body, labeled, hit,
                                 any(full[rid].get(k, False) for k in sids), late,
                                 late and bool(sids & past[rid].keys())))

        total = {k: sum(v[k] for v in summary.values()) for k in
                 ("labeled", "found", "wrong", "labeled_after", "found_after_full", "found_after_past", "wrong_after_past")}
        for subject, s in summary.items():
            print(f"- {subject[:18]}: 문제 리뷰 {s['labeled']} 중 찾음 {s['found']} · 엉뚱하게 붙음 {s['wrong']} | "
                  f"{args.cutoff} 이후 {s['labeled_after']} 중 지금 표현 {s['found_after_full']} · "
                  f"그 전 표현만 {s['found_after_past']} (표현 {s['kept'][0]}/{s['kept'][1]}) · 엉뚱 {s['wrong_after_past']}")
        print(f"total: 찾음 {total['found']}/{total['labeled']} · 엉뚱 {total['wrong']} | 이후 지금 표현 "
              f"{total['found_after_full']}/{total['labeled_after']} · 그 전 표현만 {total['found_after_past']}/"
              f"{total['labeled_after']} · 엉뚱 {total['wrong_after_past']}")
        write_sheet(rows, args.cutoff)
        print(f"sheet: {OUT.relative_to(ROOT)} ({len(rows)} rows)")
        marked = by_marks(args.cutoff)
        if marked:
            attached = {k[1]: v for k, v in marked.items() if isinstance(k, tuple) and k[0] == "붙음"}
            missed = {k[1]: v for k, v in marked.items() if isinstance(k, tuple) and k[0] == "안 붙음"}
            print(f"by marks: 붙은 것 {attached} · 안 붙은 것 {missed} | {args.cutoff} 이후 맞다 "
                  f"{marked.get('이후 맞다', 0)} 중 지금 표현 {marked.get('이후 지금 표현', 0)} · 그 전 표현만 "
                  f"{marked.get('이후 그 전 표현만', 0)} · 엉뚱(그 전 표현만) {marked.get('이후 엉뚱(그 전 표현만)', 0)}")
    finally:
        for sid, (_, cid, org) in accepted.items():
            psql((f"delete from org_knowledge_chunks where source_id='{sid}'; delete from org_knowledge_sources where id='{sid}'; "
                  if org else
                  f"delete from product_knowledge_chunks where source_id='{sid}'; "
                  f"delete from product_knowledge_sources where id='{sid}'; ")
                 + f"update knowledge_candidate set state='OPEN', source_id=null, decided_at=null, decided_by=null "
                   f"where id='{cid}';")
        print(f"removed {len(accepted)} guidance; the copy is as it was")


MARKS = ("Claude 1차", "Claude 이유", "판정", "메모")


def earlier_marks():
    """The marks a person (or Claude's first pass) already wrote, by (guidance, review text)."""
    if not OUT.exists():
        return {}
    ws = load_workbook(OUT)["근거 판정"]
    head = [c.value for c in ws[1]]
    at = {name: head.index(name) for name in ("안내", "리뷰") + MARKS if name in head}
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        out[(row[at["안내"]], row[at["리뷰"]])] = tuple(row[at[m]] if m in at else None for m in MARKS)
    return out


def by_marks(cutoff):
    """Counts by the marks in the sheet: the first-pass or the owner's 판정, whichever is there (판정 wins)."""
    ws = load_workbook(OUT)["근거 판정"]
    head = [c.value for c in ws[1]]
    ix = {h: i for i, h in enumerate(head)}
    c = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        mark = row[ix["판정"]] or row[ix["Claude 1차"]]
        if not mark:
            continue
        hit = row[ix["근거로 붙음"]] == "예"
        late = row[ix[f"{cutoff} 이후"]] == "예"
        past = row[ix["그 전 표현만으로도 붙음"]] == "예"
        c[("붙음" if hit else "안 붙음", mark)] = c.get(("붙음" if hit else "안 붙음", mark), 0) + 1
        if late and mark == "맞다":
            c["이후 맞다"] = c.get("이후 맞다", 0) + 1
            c["이후 지금 표현"] = c.get("이후 지금 표현", 0) + hit
            c["이후 그 전 표현만"] = c.get("이후 그 전 표현만", 0) + past
        if late and mark == "아니다":
            c["이후 엉뚱(그 전 표현만)"] = c.get("이후 엉뚱(그 전 표현만)", 0) + past
    return c


def write_sheet(rows, cutoff):
    marks = earlier_marks()
    book = Workbook()
    ws = book.active
    ws.title = "근거 판정"
    head = ["안내", "별점", "날짜", "리뷰", "기준 패턴이 이 문제라고 봄", "근거로 붙음", "고객 표현으로 붙음",
            f"{cutoff} 이후", "그 전 표현만으로도 붙음", "Claude 1차", "Claude 이유", "판정", "메모"]
    ws.append(head)
    for subject, rid, rating, day, body, labeled, hit, by_phrase, late, hit_past in sorted(
            rows, key=lambda r: (r[0], not r[6], r[3])):
        kept = marks.get((subject, body), (None, None, None, None))
        ws.append([subject, rating, day, body, "예" if labeled else "", "예" if hit else "", "예" if by_phrase else "",
                   "예" if late else "", "예" if hit_past else "", *kept])
    for i, width in enumerate([24, 6, 11, 70, 12, 10, 12, 10, 14, 10, 30, 10, 30]):
        ws.column_dimensions[chr(65 + i)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EEF7")
    for row in ws.iter_rows(min_row=2):
        row[3].alignment = Alignment(wrap_text=True, vertical="top")
        row[11].fill = PatternFill("solid", fgColor="FFF4C2")
    choice = DataValidation(type="list", formula1='"맞다,일부,아니다"', allow_blank=True)
    ws.add_data_validation(choice)
    choice.add(f"L2:L{ws.max_row}")
    guide = book.create_sheet("읽는 법")
    for line in [
        "한 줄 = 안내 하나와 리뷰 하나. 「판정」만 고르시면 됩니다.",
        "  맞다: 이 안내가 이 리뷰의 답글 근거로 맞다.   일부: 관련은 있지만 이 리뷰의 핵심은 아니다.   아니다: 붙으면 안 된다.",
        "「근거로 붙음」이 비어 있고 「기준 패턴이 이 문제라고 봄」이 예인 줄은 놓친 리뷰입니다 — 붙었어야 하면 「맞다」.",
        "「그 전 표현만으로도 붙음」은 표현 목록을 그날까지의 리뷰로만 만들었어도 찾았는지입니다 — 처음 보는 리뷰에서의 예상치.",
        "이 파일에는 실제 리뷰 원문이 들어 있습니다. 이 PC 밖으로 보내지 마세요.",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 120
    OUT.parent.mkdir(parents=True, exist_ok=True)
    book.save(OUT)


if __name__ == "__main__":
    main()
